import json
import os
import sys
import time
from pathlib import Path

import pytest

from planner_api.materials.parsers import MaterialParseError
from planner_api.materials.sandbox_parser import parse_material
from planner_api.materials.service import MaterialService
from planner_api.worker_sandbox import run_contained, worker_environment


def test_real_worker_denies_credentials_files_network_and_processes(tmp_path, monkeypatch):
    monkeypatch.setenv("SANDBOX_TEST_SECRET", "must-not-reach-parser")
    outside = tmp_path / "private.env"
    outside.write_text("sensitive-value")
    folder = tmp_path / "worker"
    folder.mkdir()
    run_contained([sys.executable, "-I", str(Path(__file__).with_name("sandbox_probe.py")),
        "policy", str(outside)], folder, env=worker_environment(folder), timeout_seconds=15)
    checks = json.loads((folder / "result.json").read_text())
    assert checks and all(checks.values()), checks
    assert outside.read_text() == "sensitive-value"


def test_real_os_memory_budget(tmp_path):
    run_contained([sys.executable, "-I", str(Path(__file__).with_name("sandbox_probe.py")),
        "memory"], tmp_path, env=worker_environment(tmp_path), timeout_seconds=15,
        memory_mb=512)
    assert json.loads((tmp_path / "memory.json").read_text())["limited"]


@pytest.mark.parametrize("cancel", [False, True])
def test_real_timeout_and_cancellation_kill_descendant(tmp_path, cancel):
    import psutil
    class Cancelled(Exception):
        pass
    def check():
        if cancel and (tmp_path / "child.pid").exists():
            raise Cancelled
    with pytest.raises(Cancelled if cancel else TimeoutError):
        run_contained([sys.executable, "-I", str(Path(__file__).with_name("sandbox_probe.py")),
            "tree"], tmp_path, env=worker_environment(tmp_path), timeout_seconds=2,
            check_cancelled=check)
    pid = int((tmp_path / "child.pid").read_text())
    for _ in range(40):
        if not psutil.pid_exists(pid):
            break
        time.sleep(0.05)
    assert not psutil.pid_exists(pid), "orphaned parser descendant"


@pytest.mark.parametrize("extension", [".txt", ".docx", ".pptx", ".xlsx", ".pdf"])
def test_normal_documents_parse_in_contained_worker(tmp_path, extension, pdf_bytes_factory):
    path = tmp_path / ("文档" + extension)
    if extension == ".txt":
        path.write_text("English reading practice and review", encoding="utf-8")
    elif extension == ".docx":
        from docx import Document
        doc = Document()
        doc.add_paragraph("English reading practice and review")
        doc.save(path)
    elif extension == ".pptx":
        from pptx import Presentation
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = "English reading practice"
        deck.save(path)
    elif extension == ".xlsx":
        from openpyxl import Workbook
        book = Workbook()
        book.active.append(["English reading practice", 30])
        book.save(path)
        book.close()
    else:
        path.write_bytes(pdf_bytes_factory(["English reading practice and review"]))
    result = parse_material(path, max_chars=10000)
    assert "English reading practice" in result.text


def test_parse_error_and_text_budget_remain_enforced(tmp_path):
    path = tmp_path / "large.txt"
    path.write_text("x" * 100)
    with pytest.raises(MaterialParseError) as error:
        parse_material(path, max_chars=10)
    assert error.value.code == "EXTRACTED_TEXT_TOO_LARGE"


def test_upload_rejects_parent_junction_before_any_write(tmp_path):
    from fastapi import HTTPException
    root = tmp_path / "uploads"
    outside = tmp_path / "outside"
    root.mkdir()
    outside.mkdir()
    linked = root / "goal-x"
    if os.name == "nt":
        import subprocess
        result = subprocess.run(["cmd", "/c", "mklink", "/J", str(linked), str(outside)],
                                capture_output=True)
        assert result.returncode == 0
    else:
        linked.symlink_to(outside, target_is_directory=True)
    try:
        with pytest.raises(HTTPException) as error:
            MaterialService._check_upload_destination(root, linked / "nested" / "file.txt")
        assert error.value.status_code == 422
        assert not list(outside.iterdir())
    finally:
        linked.rmdir() if os.name == "nt" else linked.unlink()


@pytest.mark.parametrize("name", ["../x.txt", "folder/.. /x.txt", "NUL.txt", "x.txt:stream",
                                 "folder./x.txt", "folder/x?.txt"])
def test_upload_rejects_windows_path_aliases(name):
    from fastapi import HTTPException
    with pytest.raises(HTTPException):
        MaterialService._normalize_uploaded_path(name)
