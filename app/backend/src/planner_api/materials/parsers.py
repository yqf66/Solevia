from __future__ import annotations

import csv
import importlib.metadata
import json
import re
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path


@dataclass(frozen=True)
class ParsedLocation:
    char_start: int
    char_end: int
    detail: dict[str, object]


@dataclass(frozen=True)
class ParsedMaterial:
    text: str
    parser_name: str
    parser_version: str
    metadata: dict[str, object]
    locations: tuple[ParsedLocation, ...] = ()


class MaterialParseError(RuntimeError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".html", ".htm", ".xml", ".json", ".csv"}
_MAX_ARCHIVE_ENTRIES = 20_000


@dataclass
class _TextBuilder:
    max_chars: int
    parts: list[str] = field(default_factory=list)
    locations: list[ParsedLocation] = field(default_factory=list)
    length: int = 0

    def append(self, text: str, detail: dict[str, object]) -> None:
        normalized = _normalize_text(text)
        if not normalized:
            return
        separator = "\n\n" if self.parts else ""
        start = self.length + len(separator)
        end = start + len(normalized)
        if end > self.max_chars:
            raise MaterialParseError(
                "EXTRACTED_TEXT_TOO_LARGE",
                "解析后的文本超过本地资料处理上限。",
            )
        if separator:
            self.parts.append(separator)
        self.parts.append(normalized)
        self.length = end
        self.locations.append(ParsedLocation(start, end, detail))

    def build(self) -> tuple[str, tuple[ParsedLocation, ...]]:
        return "".join(self.parts), tuple(self.locations)


def parse_material(
    path: Path,
    *,
    max_chars: int = 25_000_000,
    max_archive_uncompressed_bytes: int = 512 * 1024 * 1024,
    ocr_enabled: bool = False,
    ocr_max_pages: int = 100,
    ocr_dpi: int = 144,
    ocr_min_confidence: float = 0.6,
    mineru_enabled: bool = False,
    mineru_mode: str = "AUTO",
    mineru_tier: str = "standard",
    mineru_auto_threshold: int = 2,
    mineru_python: str = "",
    mineru_home: str = "",
    mineru_timeout_seconds: float = 600,
    mineru_max_pages: int = 100,
    check_cancelled=None,
) -> ParsedMaterial:
    extension = path.suffix.casefold()
    try:
        if extension == ".csv":
            builder = _TextBuilder(max_chars=max_chars)
            with path.open(encoding="utf-8", errors="replace", newline="") as stream:
                for row_number, row in enumerate(csv.reader(stream), start=1):
                    builder.append(
                        " | ".join(cell.strip() for cell in row),
                        {"kind": "csv_row", "row": row_number},
                    )
            text, locations = builder.build()
            if not text:
                raise MaterialParseError("EMPTY_DOCUMENT", "解析成功但文档没有可读取正文。")
            return ParsedMaterial(
                text=text,
                parser_name="builtin-text",
                parser_version="2",
                metadata={
                    "extension": extension,
                    "character_count": len(text),
                    "row_count": len(locations),
                    "location_count": len(locations),
                },
                locations=locations,
            )

        if extension in TEXT_EXTENSIONS:
            text = _parse_textual(path, extension, max_chars)
            normalized = _normalize_text(text)
            if not normalized:
                raise MaterialParseError("EMPTY_DOCUMENT", "解析成功但文档没有可读取正文。")
            _ensure_within_limit(normalized, max_chars)
            return ParsedMaterial(
                text=normalized,
                parser_name="builtin-text",
                parser_version="2",
                metadata={"extension": extension, "character_count": len(normalized)},
            )

        if extension in {".docx", ".pptx", ".xlsx"}:
            _check_archive_size(path, max_archive_uncompressed_bytes)

        builder = _TextBuilder(max_chars=max_chars)
        if extension == ".docx":
            metadata = _parse_docx(path, builder)
            parser_name, distribution = "python-docx", "python-docx"
        elif extension == ".pptx":
            metadata = _parse_pptx(path, builder)
            parser_name, distribution = "python-pptx", "python-pptx"
        elif extension == ".xlsx":
            metadata = _parse_xlsx(path, builder)
            parser_name, distribution = "openpyxl", "openpyxl"
        elif extension == ".pdf":
            metadata = _parse_pdf(
                path,
                builder,
                ocr_enabled=ocr_enabled,
                ocr_max_pages=ocr_max_pages,
                ocr_dpi=ocr_dpi,
                ocr_min_confidence=ocr_min_confidence,
                mineru_enabled=mineru_enabled,
                mineru_mode=mineru_mode,
                mineru_tier=mineru_tier,
                mineru_auto_threshold=mineru_auto_threshold,
                mineru_python=mineru_python,
                mineru_home=mineru_home,
                mineru_timeout_seconds=mineru_timeout_seconds,
                mineru_max_pages=mineru_max_pages,
                check_cancelled=check_cancelled,
            )
            parser_name, distribution = "pdfplumber", "pdfplumber"
        elif extension in {".xls"}:
            text = _parse_markitdown(path)
            _ensure_within_limit(text, max_chars)
            normalized = _normalize_text(text)
            return ParsedMaterial(
                text=normalized,
                parser_name="markitdown",
                parser_version=_package_version("markitdown", "unknown"),
                metadata={"extension": extension, "character_count": len(normalized)},
            )
        else:
            raise MaterialParseError(
                "UNSUPPORTED_FORMAT", f"暂不支持 {extension or '无扩展名'} 文件"
            )

        text, locations = builder.build()
        enhanced_pages = metadata.pop("_mineru_pages", None)
        if enhanced_pages:
            enhanced_builder = _TextBuilder(max_chars=max_chars)
            for page in enhanced_pages:
                enhanced_builder.append(
                    page["text"], {"kind": "pdf_page", "page": page["page"], "extraction": "MinerU"}
                )
            text, locations = enhanced_builder.build()
            parser_name, distribution = "mineru", "mineru"
        if not text and metadata.get("unread_pages"):
            raise MaterialParseError(
                "OCR_REQUIRED", "扫描页没有可读取正文；请启用 OCR 并检查本地模型及解析警告。"
            )
        if not text:
            raise MaterialParseError("EMPTY_DOCUMENT", "解析成功但文档没有可读取正文。")
        return ParsedMaterial(
            text=text,
            parser_name=parser_name,
            parser_version=metadata.get("mineru_version")
            or _package_version(distribution, "unknown"),
            metadata={
                "extension": extension,
                "character_count": len(text),
                "location_count": len(locations),
                **metadata,
            },
            locations=locations,
        )
    except MaterialParseError:
        raise
    except Exception as error:
        if check_cancelled:
            check_cancelled()
        raise MaterialParseError("PARSE_FAILED", f"资料解析失败：{error}") from error


def _parse_textual(path: Path, extension: str, max_chars: int) -> str:
    if extension in {".html", ".htm", ".xml"}:
        return _markup_file_to_text(path, max_chars)
    with path.open(encoding="utf-8", errors="replace") as stream:
        raw = stream.read(max_chars + 1)
    _ensure_within_limit(raw, max_chars)
    if extension == ".json":
        try:
            raw = json.dumps(json.loads(raw), ensure_ascii=False, indent=2)
        except json.JSONDecodeError:
            pass
    return raw


class _VisibleHTMLText(HTMLParser):
    def __init__(self, max_chars: int) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._ignored_tags: list[str] = []
        self.max_chars = max_chars
        self.visible_chars = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.casefold() in {"script", "style", "noscript"}:
            self._ignored_tags.append(tag.casefold())
        elif tag.casefold() in {"p", "div", "li", "tr", "br", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        tag = tag.casefold()
        if self._ignored_tags and tag == self._ignored_tags[-1]:
            self._ignored_tags.pop()
        elif tag in {"p", "div", "li", "tr", "br", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._ignored_tags:
            self.visible_chars += len(data)
            if self.visible_chars > self.max_chars:
                raise MaterialParseError(
                    "EXTRACTED_TEXT_TOO_LARGE",
                    "解析后的文本超过本地资料处理上限。",
                )
            self.parts.append(data)


def _markup_file_to_text(path: Path, max_chars: int) -> str:
    parser = _VisibleHTMLText(max_chars)
    with path.open(encoding="utf-8", errors="replace") as stream:
        while block := stream.read(64 * 1024):
            parser.feed(block)
    parser.close()
    return "".join(parser.parts)


def _parse_docx(path: Path, builder: _TextBuilder) -> dict[str, object]:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as error:
        raise MaterialParseError(
            "PARSER_UNAVAILABLE", "DOCX 解析需要 python-docx，请安装 materials 依赖。"
        ) from error

    document = Document(str(path))
    heading_path: list[str] = []
    paragraph_count = 0
    table_count = 0
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            paragraph_count += 1
            text = block.text.strip()
            if not text:
                continue
            style_name = block.style.name if block.style is not None else ""
            level = _heading_level(style_name)
            if level is not None:
                heading_path = heading_path[: level - 1] + [text]
            detail = {
                "kind": "docx_paragraph",
                "paragraph": paragraph_count,
                "heading_path": heading_path.copy(),
            }
            builder.append(text, detail)
        elif isinstance(block, Table):
            table_count += 1
            for row_index, row in enumerate(block.rows, start=1):
                cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
                while cells and not cells[-1]:
                    cells.pop()
                if any(cells):
                    builder.append(
                        " | ".join(cells),
                        {
                            "kind": "docx_table_row",
                            "table": table_count,
                            "row": row_index,
                            "heading_path": heading_path.copy(),
                        },
                    )
    return {"paragraph_count": paragraph_count, "table_count": table_count}


def _parse_pptx(path: Path, builder: _TextBuilder) -> dict[str, object]:
    try:
        from pptx import Presentation
    except ImportError as error:
        raise MaterialParseError(
            "PARSER_UNAVAILABLE", "PPTX 解析需要 python-pptx，请安装 materials 依赖。"
        ) from error

    presentation = Presentation(str(path))
    text_shape_count = 0
    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_detail = {"kind": "pptx_slide", "slide": slide_number}
        for shape_index, shape in enumerate(slide.shapes, start=1):
            if not hasattr(shape, "text"):
                continue
            value = shape.text.strip()
            if not value:
                continue
            text_shape_count += 1
            builder.append(
                value,
                {
                    **slide_detail,
                    "kind": "pptx_text_shape",
                    "shape": shape_index,
                    "shape_name": shape.name[:120],
                },
            )
    return {"slide_count": len(presentation.slides), "text_shape_count": text_shape_count}


def _parse_xlsx(path: Path, builder: _TextBuilder) -> dict[str, object]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise MaterialParseError(
            "PARSER_UNAVAILABLE", "XLSX 解析需要 openpyxl，请安装 materials 依赖。"
        ) from error

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    row_count = 0
    try:
        for sheet in workbook.worksheets:
            has_content = False
            for row_number, row in enumerate(sheet.iter_rows(values_only=False), start=1):
                values: list[str] = []
                cell_refs: list[str] = []
                for cell in row:
                    if cell.value is None:
                        continue
                    value = str(cell.value).strip()
                    if value:
                        values.append(value)
                        if len(cell_refs) < 32:
                            cell_refs.append(cell.coordinate)
                if values:
                    if not has_content:
                        builder.append(
                            f"## Sheet: {sheet.title}",
                            {"kind": "xlsx_sheet", "sheet": sheet.title},
                        )
                        has_content = True
                    row_count += 1
                    builder.append(
                        " | ".join(values),
                        {
                            "kind": "xlsx_row",
                            "sheet": sheet.title,
                            "row_start": row_number,
                            "row_end": row_number,
                            "cell_refs": cell_refs,
                        },
                    )
    finally:
        workbook.close()
    return {"sheet_count": len(workbook.worksheets), "nonempty_row_count": row_count}


def _parse_pdf(
    path: Path,
    builder: _TextBuilder,
    *,
    ocr_enabled=False,
    ocr_max_pages=100,
    ocr_dpi=144,
    ocr_min_confidence=0.6,
    mineru_enabled=False,
    mineru_mode="AUTO",
    mineru_tier="standard",
    mineru_auto_threshold=2,
    mineru_python="",
    mineru_home="",
    mineru_timeout_seconds=600,
    mineru_max_pages=100,
    check_cancelled=None,
) -> dict[str, object]:
    try:
        import pdfplumber
    except ImportError as error:
        raise MaterialParseError(
            "PARSER_UNAVAILABLE", "PDF 解析需要 pdfplumber，请安装 materials 依赖。"
        ) from error

    warnings = []
    ocr_pages = []
    unread_pages = []
    ocr_attempts = 0
    complexity_score = 0
    complexity_reasons: set[str] = set()
    page_count = 0
    with pdfplumber.open(path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            page_count += 1
            if check_cancelled:
                check_cancelled()
            text = page.extract_text() or ""
            tables = page.extract_tables() or []
            page_score, page_reasons = (
                _pdf_complexity_signals(page, text, tables) if mineru_enabled else (0, set())
            )
            # One complex page suffices; repeated simple pages must not inflate the score.
            complexity_score = max(complexity_score, page_score)
            complexity_reasons.update(page_reasons)
            if text.strip():
                builder.append(text, {"kind": "pdf_page", "page": page_number})
            has_scan = any(
                image.get("width", 0) * image.get("height", 0) >= page.width * page.height * 0.4
                for image in page.images
            )
            if has_scan and len(text.strip()) < 40:
                if ocr_enabled and ocr_attempts < ocr_max_pages:
                    ocr_attempts += 1
                    try:
                        from .ocr import recognize_page

                        scanned_text, confidence, rejected_lines = recognize_page(
                            page, dpi=ocr_dpi, min_confidence=ocr_min_confidence
                        )
                        if scanned_text.strip():
                            # Replace sparse text (e.g. an embedded page number) to avoid double counting.
                            if text.strip():
                                location = builder.locations.pop()
                                builder.parts.pop()
                                if builder.parts and builder.parts[-1] == "\n\n":
                                    builder.parts.pop()
                                builder.length = location.char_start - (2 if builder.parts else 0)
                            builder.append(
                                scanned_text,
                                {
                                    "kind": "pdf_page",
                                    "page": page_number,
                                    "extraction": "OCR",
                                    "confidence": round(confidence, 4),
                                },
                            )
                            ocr_pages.append(page_number)
                            if rejected_lines:
                                unread_pages.append(page_number)
                                warnings.append(
                                    f"第 {page_number} 页有 {rejected_lines} 行低置信度文字未纳入正文，请核对原文件"
                                )
                        else:
                            unread_pages.append(page_number)
                    except MaterialParseError:
                        raise
                    except Exception as error:
                        unread_pages.append(page_number)
                        warnings.append(f"第 {page_number} 页 OCR 未完成：{str(error)[:200]}")
                else:
                    unread_pages.append(page_number)
                    warnings.append(f"第 {page_number} 页为扫描内容，OCR 未启用或超过页数预算")
            for table_number, table in enumerate(tables, start=1):
                for row_number, row in enumerate(table, start=1):
                    cells = [str(cell).strip() for cell in row if cell is not None]
                    if any(cells):
                        builder.append(
                            " | ".join(cells),
                            {
                                "kind": "pdf_table_row",
                                "page": page_number,
                                "table": table_number,
                                "row": row_number,
                            },
                        )
            page.close()
        should_use_mineru = mineru_enabled and (
            mineru_mode == "FORCE" or complexity_score >= mineru_auto_threshold
        )
        metadata: dict[str, object] = {
            "page_count": page_count,
            "ocr_pages": ocr_pages,
            "unread_pages": unread_pages,
            "warnings": warnings,
            "coverage": "PARTIAL" if unread_pages else "COMPLETE",
            "complexity_score": complexity_score,
            "complexity_reasons": sorted(complexity_reasons),
            "mineru_attempted": should_use_mineru,
            "mineru_used": False,
        }
        if should_use_mineru:
            try:
                if check_cancelled:
                    check_cancelled()
                if page_count > mineru_max_pages:
                    raise RuntimeError(f"MinerU 页数超过预算（{page_count}/{mineru_max_pages}）")
                from .mineru_adapter import parse_pdf as parse_with_mineru

                enhanced = parse_with_mineru(
                    path,
                    tier=mineru_tier,
                    python=mineru_python,
                    home=mineru_home,
                    timeout_seconds=mineru_timeout_seconds,
                    max_chars=builder.max_chars,
                    check_cancelled=check_cancelled,
                )
                enhanced_pages = enhanced["pages"]
                if [p["page"] for p in enhanced_pages] != list(range(1, page_count + 1)):
                    raise RuntimeError("MinerU 页码覆盖不完整，保留基础解析结果")
                if not any(p["text"].strip() for p in enhanced_pages):
                    raise RuntimeError("MinerU 没有返回可读取正文")
                # Validate budgets before replacing the working fallback.
                enhanced_builder = _TextBuilder(max_chars=builder.max_chars)
                for item in enhanced_pages:
                    enhanced_builder.append(
                        item["text"],
                        {"kind": "pdf_page", "page": item["page"], "extraction": "MinerU"},
                    )
                metadata["_mineru_pages"] = enhanced_pages
                metadata["mineru_used"] = True
                metadata["mineru_version"] = enhanced.get("version", "unknown")
                metadata["parser_selection"] = "mineru"
                metadata["basic_warnings"] = warnings
                metadata["warnings"] = []
                metadata["unread_pages"] = [
                    p["page"] for p in enhanced_pages if not p["text"].strip()
                ]
                metadata["coverage"] = "PARTIAL" if metadata["unread_pages"] else "COMPLETE"
            except Exception as error:
                if check_cancelled:
                    check_cancelled()
                metadata["warnings"] = [
                    *warnings,
                    f"复杂 PDF 已尝试 MinerU，但未完成，已回退基础解析：{str(error)[:240]}",
                ]
                metadata["mineru_warning"] = str(error)[:240]
        return metadata


def _pdf_complexity_signals(page, text: str, tables=None) -> tuple[int, set[str]]:
    score = 0
    reasons: set[str] = set()
    images = getattr(page, "images", ()) or ()
    page_area = float(getattr(page, "width", 0) or 0) * float(getattr(page, "height", 0) or 0)
    if page_area and any(
        float(image.get("width", 0)) * float(image.get("height", 0)) >= page_area * 0.4
        for image in images
    ):
        score += 1
        reasons.add("large_image_or_scan")
    if tables is None:
        tables = page.extract_tables() or []
    if tables:
        # Merged cells or larger grids are more likely to need layout recovery.
        complex_table = any(
            len(table) >= 4
            and (
                max((len(row) for row in table), default=0) >= 3
                or any(cell is None for row in table for cell in row)
            )
            for table in tables
        )
        score += 2 if complex_table else 1
        reasons.add("table_layout")
    try:
        words = page.extract_words() or []
    except Exception:
        words = []
    if len(words) >= 12 and getattr(page, "width", 0):
        width = float(page.width)
        gutter = width * 0.04
        # Require a vertical gutter; full-width paragraphs have words crossing it.
        left = [word for word in words if float(word.get("x1", 0)) <= width * 0.5 - gutter]
        right = [word for word in words if float(word.get("x0", 0)) >= width * 0.5 + gutter]
        crossing = [
            word
            for word in words
            if float(word.get("x0", 0)) < width * 0.5 + gutter
            and float(word.get("x1", 0)) > width * 0.5 - gutter
        ]
        aligned = len(
            {round(float(word.get("top", 0)) / 8) for word in left}
            & {round(float(word.get("top", 0)) / 8) for word in right}
        )
        if (
            len(left) >= 4
            and len(right) >= 4
            and aligned >= 3
            and len(crossing) <= len(words) * 0.05
        ):
            score += 2
            reasons.add("multi_column_layout")
    if re.search(r"[∑∫√≈≠≤≥∞∂∇]|\\(?:frac|sum|int|sqrt)|\b[A-Za-z]\s*[=<>]\s*[A-Za-z0-9]", text):
        score += 2
        reasons.add("formula_or_math")
    return score, reasons


def _parse_markitdown(path: Path) -> str:
    try:
        from markitdown import MarkItDown
    except ImportError as error:
        raise MaterialParseError(
            "PARSER_UNAVAILABLE",
            "该文件需要 MarkItDown，但当前环境未安装对应解析依赖。",
        ) from error
    result = MarkItDown().convert_local(path)
    text = getattr(result, "text_content", "") or ""
    if not text.strip():
        raise MaterialParseError("EMPTY_DOCUMENT", "解析成功但文档没有可读取正文。")
    return text


def _check_archive_size(path: Path, max_uncompressed_bytes: int) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
    except zipfile.BadZipFile as error:
        raise MaterialParseError("INVALID_ARCHIVE", "Office 文档不是有效的 ZIP 文档。") from error
    if len(entries) > _MAX_ARCHIVE_ENTRIES:
        raise MaterialParseError("ARCHIVE_TOO_COMPLEX", "Office 文档包含过多内部文件。")
    expanded_bytes = sum(entry.file_size for entry in entries)
    if expanded_bytes > max_uncompressed_bytes:
        raise MaterialParseError("ARCHIVE_TOO_LARGE", "Office 文档解压后的内容超过本地安全上限。")


def _heading_level(style_name: str) -> int | None:
    match = re.search(r"(?:heading|标题)\s*(\d+)", style_name, flags=re.I)
    if match is None:
        return None
    return min(max(int(match.group(1)), 1), 6)


def _ensure_within_limit(text: str, max_chars: int) -> None:
    if len(text) > max_chars:
        raise MaterialParseError("EXTRACTED_TEXT_TOO_LARGE", "解析后的文本超过本地资料处理上限。")


def _package_version(distribution: str, fallback: str) -> str:
    try:
        return importlib.metadata.version(distribution)
    except importlib.metadata.PackageNotFoundError:
        return fallback


def _normalize_text(value: str) -> str:
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in value.splitlines()]
    compact: list[str] = []
    blank = False
    for line in lines:
        if not line:
            if compact and not blank:
                compact.append("")
            blank = True
            continue
        compact.append(line)
        blank = False
    return "\n".join(compact).strip()
